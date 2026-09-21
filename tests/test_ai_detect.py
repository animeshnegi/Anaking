"""
Tests for the AI-generated / pasted answer check on open-text questions.

    python3 -m pytest tests/test_ai_detect.py

Covers the detector itself (core.ai_detect), the QC flags it feeds (core.qc), the live
endpoint the survey calls while a respondent types (/api/check_text), the Admin review
queue (/api/admin/verbatims) and the export columns / sheet.
"""

import csv
import io

from core.ai_detect import (VERDICT_HUMAN, VERDICT_LIKELY, VERDICT_POSSIBLE, ai_settings,
                            duplicate_verbatims, proofread, score_text, text_fields)
from core.qc import qc_flags

# Real-sounding verbatims: contractions, first person, concrete numbers, rough mechanics.
HUMAN = [
    "honestly the once-a-week shot is the big one for us. we run infusion 4 days a week and "
    "chairs are full by 9am, so anything that frees up a chair gets used. the AE rate worries "
    "me a bit - 30% grade 3+ is not nothing and i'd want the breakdown by line of therapy "
    "before putting it in a frail 78 year old. probably 10-15% of my NSCLC patients in year 1.",
    "Cost. It's always cost. The drug works, I believe the PFS data, but my patients are on "
    "fixed incomes and the copay assistance only lasts 12 months. After that they stop. I've "
    "watched it happen twice now with the other agent.",
    "I don't love the subcutaneous route - nurses hate the volume and patients complain about "
    "the site reactions for two days. But the time saving is real. We'd use it for maintenance "
    "patients mostly, not first line.",
    "not sure yet. need to see the full OS data at ASCO before i change anything. right now "
    "pembro+chemo is what i know and the payers cover it without a fight.",
]

# The same questions answered by a chatbot: hallmark vocabulary, even rhythm, no
# contractions, no first person, markdown, refusal phrasing.
AI = [
    "Oncologists play a crucial role in navigating the complex landscape of advanced NSCLC "
    "treatment. Moreover, the safety profile appears robust, with grade 3+ adverse events "
    "manageable in most cases. Furthermore, once-weekly subcutaneous administration offers a "
    "compelling convenience advantage over intravenous alternatives, which could streamline "
    "clinic workflows. Additionally, the biomarker-driven approach underscores the importance "
    "of comprehensive genomic profiling. In conclusion, I would anticipate meaningful uptake "
    "among community practices where infusion capacity is constrained.",
    "There are several key factors that would influence my willingness to prescribe this "
    "product. Firstly, the efficacy data must demonstrate a clear progression-free survival "
    "benefit. Secondly, the safety profile should be manageable in a real-world population. "
    "Thirdly, reimbursement and access considerations play a vital role. Finally, convenience "
    "of administration cannot be overlooked. Overall, a multifaceted approach is essential to "
    "successful adoption in clinical practice.",
    "As an AI language model, I do not have personal medical opinions or clinical experience. "
    "However, I can provide a general overview of the factors that oncologists typically "
    "consider when evaluating a new therapeutic option in advanced non-small cell lung cancer.",
    "**Key considerations**\n\n- Efficacy: robust PFS benefit\n- Safety: manageable grade 3+ "
    "events\n- Access: reimbursement pathways\n\nIn summary, this product represents a "
    "promising addition to the treatment landscape, offering a nuanced balance of convenience "
    "and clinical value for oncologists and patients alike.",
]

STUDY = {
    "sections": [{"id": "S1", "title": "A"}],
    "questions": [
        {"id": "Q1", "section": "S1", "type": "single_select", "stem": "Setting",
         "options": [{"code": 1, "label": "Community"},
                     {"code": 2, "label": "Other", "other": True}]},
        {"id": "Q2", "section": "S1", "type": "open_text", "stem": "Why?"},
        {"id": "Q3", "section": "S1", "type": "open_text", "stem": "Anything else?",
         "ai_check": False},
    ],
    "qc": {"min_seconds": 10, "verbatim_qs": ["Q2"]},
}


# ---------------------------------------------------------------- the detector
def test_human_answers_read_as_human():
    for text in HUMAN:
        res = score_text(text)
        assert res["verdict"] == VERDICT_HUMAN, (text[:50], res["score"], res["signals"])
        assert res["score"] < 35


def test_ai_answers_are_flagged():
    for text in AI:
        res = score_text(text)
        assert res["verdict"] == VERDICT_LIKELY, (text[:50], res["score"], res["signals"])
        assert res["score"] >= 60
        assert res["signals"], "a flag must always carry the evidence behind it"


def test_every_scored_answer_explains_itself():
    for text in AI + HUMAN:
        res = score_text(text)
        assert 0 <= res["score"] <= 100
        for sig in res["signals"]:
            assert sig["label"] and isinstance(sig["weight"], (int, float))


def test_short_answers_are_not_judged():
    res = score_text("It is good.")
    assert res["verdict"] == "too_short" and res["scored"] is False and res["score"] == 0


def test_ai_disclaimer_is_flagged_even_in_a_short_answer():
    res = score_text("As an AI language model, I cannot answer that.")
    assert res["verdict"] == VERDICT_LIKELY
    assert any(s["key"] == "disclaimer" for s in res["signals"])


def test_pasted_telemetry_pushes_an_answer_over_the_line():
    text = ("There are several considerations here. Firstly, efficacy. Secondly, safety. "
            "Thirdly, access and reimbursement. Overall, a multifaceted approach matters.")
    plain = score_text(text)
    pasted = score_text(text, {"keystrokes": 0, "pastes": 1, "pasted_chars": len(text),
                               "input_events": 1, "typed_ms": 200, "blur_ms": 30000})
    assert pasted["score"] > plain["score"]
    assert pasted["verdict"] == VERDICT_LIKELY
    assert {s["key"] for s in pasted["signals"]} >= {"pasted_all", "no_keystrokes", "too_fast"}


def test_typing_telemetry_keeps_a_human_answer_clean():
    res = score_text(HUMAN[0], {"keystrokes": 330, "pastes": 0, "pasted_chars": 0,
                                "input_events": 240, "typed_ms": 60000, "blur_ms": 0})
    assert res["verdict"] == VERDICT_HUMAN and res["score"] == 0
    assert any(s["key"] == "typed_normally" for s in res["signals"])


def test_thresholds_are_configurable_per_study():
    default = score_text(AI[3])
    assert default["verdict"] == VERDICT_LIKELY
    strict = score_text(AI[3], None, {"flag_at": 95, "warn_at": 90})
    assert strict["verdict"] == VERDICT_HUMAN          # same text, higher bar
    lenient = score_text(AI[3], None, {"flag_at": 10, "warn_at": 5})
    assert lenient["verdict"] == VERDICT_LIKELY
    assert strict["score"] == default["score"] == lenient["score"]   # only the bar moves


def test_proofreading_notes():
    notes = {n["key"] for n in proofread("The the answer is good **overall** and also N/A")}
    assert {"doubled_word", "markdown"} <= notes
    assert proofread("") == []
    assert proofread("We would use it for maintenance patients, about 10 a month.") == []
    assert any(n["key"] == "run_on" for n in proofread("and " * 60 + "more"))


# ---------------------------------------------------------------- study wiring
def test_text_fields_covers_every_free_text_box():
    assert text_fields(STUDY) == [("Q1", "other_text"), ("Q2", "_"), ("Q3", "_")]


def test_ai_settings_can_be_turned_off_per_question():
    assert ai_settings(STUDY, "Q2")["enabled"] is True
    assert ai_settings(STUDY, "Q2")["action"] == "confirm"       # default
    assert ai_settings(STUDY, "Q3")["enabled"] is False          # question opts out
    study_off = {"qc": {"ai": {"enabled": False}}}
    assert ai_settings(study_off, "Q2")["enabled"] is False
    warn_only = {"questions": [{"id": "Q2", "ai_action": "warn"}], "qc": {}}
    assert ai_settings(warn_only, "Q2")["action"] == "warn"


def test_duplicate_verbatims_are_matched_across_respondents():
    shared = AI[0]
    recs = [{"respondent_code": "R001", "answers": {"Q2": {"_": shared}}},
            {"respondent_code": "R002", "answers": {"Q2": {"_": shared.upper() + "  "}}},
            {"respondent_code": "R003", "answers": {"Q2": {"_": HUMAN[0]}}}]
    dups = duplicate_verbatims(recs, STUDY)
    assert sorted(dups[("R001", "Q2")]) == ["R002"]
    assert dups[("R002", "Q2")] == ["R001"]
    assert ("R003", "Q2") not in dups


# ---------------------------------------------------------------- QC flags
def test_qc_flags_raise_ai_flags_for_a_generated_answer():
    answers = {"Q2": {"_": AI[0], "_meta": {"keystrokes": 0, "pasted_chars": len(AI[0]),
                                            "pastes": 1, "typed_ms": 400}}}
    out = qc_flags(answers, 900, STUDY)
    assert "ai_generated_Q2" in out["flags"]
    assert "ai_generated_verbatim" in out["flags"]          # respondent-level roll-up
    assert out["clean"] is False
    assert out["ai"]["Q2"]["verdict"] == VERDICT_LIKELY
    assert out["ai"]["Q2"]["signals"]                       # evidence is carried through


def test_the_review_band_is_flagged_as_suspect_not_generated():
    formal = ("The main hesitation would be the safety profile and the cost. Reimbursement is "
              "always uncertain for a new agent, and my patients struggle with copays. I would "
              "want more data before changing practice, but the convenience is appealing for "
              "maintenance patients overall.")
    assert score_text(formal)["verdict"] == VERDICT_HUMAN       # just under the default bar
    strict = {"warn_at": 30, "flag_at": 80}
    assert score_text(formal, None, strict)["verdict"] == VERDICT_POSSIBLE
    cfg = {"questions": [{"id": "Q9", "type": "open_text", "stem": "x"}], "qc": {"ai": strict}}
    flags = qc_flags({"Q9": {"_": formal}}, 900, cfg)["flags"]
    assert "ai_suspect_Q9" in flags
    assert "ai_generated_verbatim" not in flags     # only "likely" trips the roll-up


def test_qc_flags_stay_clean_for_a_typed_human_answer():
    answers = {"Q2": {"_": HUMAN[0], "_meta": {"keystrokes": 330, "pasted_chars": 0,
                                               "input_events": 240, "typed_ms": 60000}}}
    out = qc_flags(answers, 900, STUDY)
    assert out["clean"] is True and out["flags"] == []
    assert out["ai"]["Q2"]["verdict"] == VERDICT_HUMAN


def test_qc_flags_respect_a_question_that_opts_out():
    answers = {"Q3": {"_": AI[0]}}
    assert qc_flags(answers, 900, STUDY)["ai"] == {}        # Q3 has ai_check: False


def test_qc_flags_check_every_open_text_question_by_default():
    cfg = {"questions": [{"id": "Q9", "type": "open_text", "stem": "x"}]}
    out = qc_flags({"Q9": {"_": AI[1]}}, 900, cfg)          # Q9 is in no verbatim_qs list
    assert "ai_generated_Q9" in out["flags"]
    limited = {"questions": [{"id": "Q9", "type": "open_text", "stem": "x"}],
               "qc": {"check_all_text": False, "verbatim_qs": []}}
    assert qc_flags({"Q9": {"_": AI[1]}}, 900, limited)["ai"] == {}


# ---------------------------------------------------------------- HTTP surface
def test_check_text_endpoint_scores_a_live_answer(client):
    r = client.post("/api/check_text", json={"study": "beacon", "qid": "Q20b",
                                             "text": AI[0]})
    body = r.get_json()
    assert r.status_code == 200 and body["verdict"] == VERDICT_LIKELY
    assert body["score"] >= 60 and body["action"] == "confirm"
    assert body["signals"] and isinstance(body["proofread"], list)

    r = client.post("/api/check_text", json={"study": "beacon", "qid": "Q20b",
                                             "text": HUMAN[1]})
    assert r.get_json()["verdict"] == VERDICT_HUMAN

    r = client.post("/api/check_text", json={"study": "beacon", "qid": "Q20b",
                                             "text": "fine"})
    assert r.get_json()["verdict"] == "too_short"


def test_check_text_uses_paste_telemetry_from_the_client(client):
    text = ("There are several considerations here. Firstly, efficacy. Secondly, safety. "
            "Thirdly, access and reimbursement. Overall, a multifaceted approach matters.")
    typed = client.post("/api/check_text", json={"qid": "Q20b", "text": text, "meta": {
        "keystrokes": 150, "pasted_chars": 0, "input_events": 120, "typed_ms": 25000}})
    pasted = client.post("/api/check_text", json={"qid": "Q20b", "text": text, "meta": {
        "keystrokes": 0, "pasted_chars": len(text), "pastes": 1, "input_events": 1,
        "typed_ms": 150}})
    assert pasted.get_json()["score"] > typed.get_json()["score"]
    assert pasted.get_json()["verdict"] == VERDICT_LIKELY


def test_spec_exposes_the_study_ai_settings(client):
    spec = client.get("/api/spec").get_json()
    assert spec["ai_check"]["action"] == "confirm"
    assert spec["ai_check"]["flag_at"] == 60


AI_STUDY = {
    "sections": [{"id": "S1", "title": "A"}],
    "questions": [
        {"id": "Q1", "section": "S1", "type": "open_text", "stem": "Why would you hesitate?"},
        {"id": "Q2", "section": "S1", "type": "open_text", "stem": "Anything else?"},
    ],
    "qc": {"min_seconds": 10},
}


def _fielded_pair(client):
    """Two respondents: one pasted a chatbot answer, one typed their own."""
    client.post("/api/studio/save", json={"title": "AI Field", "cfg": AI_STUDY})
    client.post("/api/studio/status", json={"slug": "ai-field", "status": "live"})
    out = {}
    for code, text, meta in [
        ("ai", AI[0], {"keystrokes": 0, "pastes": 1, "pasted_chars": len(AI[0]),
                       "input_events": 1, "typed_ms": 300, "blur_ms": 20000}),
        ("human", HUMAN[0], {"keystrokes": 330, "pastes": 0, "pasted_chars": 0,
                             "input_events": 240, "typed_ms": 60000}),
    ]:
        s = client.post("/api/start", json={"study": "ai-field"}).get_json()
        client.post("/api/save", json={"session_id": s["session_id"], "elapsed_seconds": 600,
                                       "answers": {"Q1": {"_": text, "_meta": meta,
                                                          "_ai": {"score": 70,
                                                                  "verdict": "likely_ai",
                                                                  "ack": code == "human"}}}})
        done = client.post("/api/submit", json={"session_id": s["session_id"],
                                                "elapsed_seconds": 600}).get_json()
        out[code] = {"sid": s["session_id"], "code": done["respondent_code"],
                     "flags": done["flags"]}
    return out


def test_ai_flag_lands_on_the_record_at_submit(client):
    out = _fielded_pair(client)
    assert "ai_generated_Q1" in out["ai"]["flags"]
    assert "ai_generated_verbatim" in out["ai"]["flags"]
    assert out["human"]["flags"] == []


def test_telemetry_and_verdict_survive_the_round_trip(client):
    out = _fielded_pair(client)
    sid = out["human"]["sid"]
    p = client.get("/api/progress", query_string={"sid": sid}).get_json()
    assert p["answers"]["Q1"]["_meta"]["keystrokes"] == 330      # decoded back to a dict
    assert p["answers"]["Q1"]["_ai"]["ack"] is True


def test_admin_dashboard_rolls_up_the_ai_check(client):
    _fielded_pair(client)
    d = client.get("/api/admin/data", query_string={"study": "ai-field"}).get_json()
    ai = d["ai_text"]
    assert ai["answers_scored"] == 2 and ai["likely_ai"] == 1
    assert ai["respondents_flagged"] == 1 and ai["confirmed_own_words"] == 1
    assert [q["id"] for q in ai["per_question"]][0] == "Q1"
    assert any("ai_generated_verbatim" in row["flags"] for row in d["qc_flagged"])


def test_admin_verbatim_queue_is_the_proofreading_list(client):
    _fielded_pair(client)
    d = client.get("/api/admin/verbatims", query_string={"study": "ai-field"}).get_json()
    assert d["summary"]["likely_ai"] == 1
    first = d["rows"][0]                                     # worst first
    assert first["verdict"] == VERDICT_LIKELY and first["score"] >= 60
    assert first["keystrokes"] == 0 and first["pasted_pct"] > 90
    assert first["signals"] and "AI-typical vocabulary" in "; ".join(first["signals"])
    assert d["rows"][-1]["verdict"] == VERDICT_HUMAN

    only = client.get("/api/admin/verbatims",
                      query_string={"study": "ai-field", "verdict": "flagged"}).get_json()
    assert [r["verdict"] for r in only["rows"]] == [VERDICT_LIKELY]
    by_q = client.get("/api/admin/verbatims",
                      query_string={"study": "ai-field", "qid": "Q2"}).get_json()
    assert by_q["rows"] == []
    search = client.get("/api/admin/verbatims",
                        query_string={"study": "ai-field", "q": "chairs are full"}).get_json()
    assert len(search["rows"]) == 1 and search["rows"][0]["verdict"] == VERDICT_HUMAN


def test_duplicate_answers_are_named_in_the_queue(client):
    client.post("/api/studio/save", json={"title": "Dup Field", "cfg": AI_STUDY})
    client.post("/api/studio/status", json={"slug": "dup-field", "status": "live"})
    for _ in range(2):
        s = client.post("/api/start", json={"study": "dup-field"}).get_json()
        client.post("/api/save", json={"session_id": s["session_id"], "elapsed_seconds": 600,
                                       "answers": {"Q1": {"_": AI[0]}}})
        client.post("/api/submit", json={"session_id": s["session_id"],
                                         "elapsed_seconds": 600})
    rows = client.get("/api/admin/verbatims",
                      query_string={"study": "dup-field"}).get_json()["rows"]
    assert rows[0]["duplicate_of"] and rows[1]["duplicate_of"]


def test_export_carries_the_ai_columns_and_a_review_sheet(client):
    _fielded_pair(client)
    txt = client.get("/admin/export.csv",
                     query_string={"study": "ai-field"}).data.decode()
    rows = list(csv.DictReader(io.StringIO(txt)))
    flagged = next(r for r in rows if "ai_generated_verbatim" in r["qc_flags"])
    assert flagged["Q1_ai_verdict"] == VERDICT_LIKELY
    assert int(flagged["Q1_ai_score"]) >= 60
    assert flagged["Q1_keystrokes"] == "0"
    assert flagged["Q1_pasted_chars"] == str(len(AI[0]))
    assert flagged["ai_generated_answers"] == "1"
    clean = next(r for r in rows if not r["qc_flags"])
    assert clean["Q1_ai_verdict"] == VERDICT_HUMAN
    assert clean["Q1_ai_confirmed_own_words"] == "yes"

    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(
        client.get("/admin/export.xlsx", query_string={"study": "ai-field"}).data))
    assert "Verbatim AI check" in wb.sheetnames
    sheet = [tuple(r) for r in wb["Verbatim AI check"].iter_rows(values_only=True)]
    head = sheet[0]
    assert {"ai_score", "verdict", "pasted_chars", "duplicate_of", "evidence"} <= set(head)
    assert sheet[1][head.index("verdict")] == VERDICT_LIKELY
    summary = {r[0]: r[1] for r in wb["Field summary"].iter_rows(values_only=True)}
    assert summary["Written answers flagged as AI-generated"] == 1
